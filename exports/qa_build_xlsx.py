import sys, io, json, datetime, collections
sys.path.insert(0, 'exports')
from qa_core import task_label, TASK, VID, USER, BYTASK
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

R = json.load(io.open('exports/qa_results.json', encoding='utf-8'))
TODAY = datetime.date.today().isoformat()

F = 'Arial'
HEAD = PatternFill('solid', fgColor='1F3864')
SUB = PatternFill('solid', fgColor='D9E1F2')
INPUT = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HAND = {'L': 'Left', 'R': 'Right', 'LR': 'Both'}


def tc(s):
    s = max(0.0, float(s)); m = int(s // 60)
    return '%02d:%05.2f' % (m, s - 60 * m)


def hand_name(h):
    return '/'.join(HAND.get(p, p) for p in str(h).split('/'))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BOX
    ws.row_dimensions[row].height = 32


def write_table(ws, headers, rows, widths, wrap_cols=(), input_cols=(), num_formats=None):
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for r in rows:
        ws.append(r)
    num_formats = num_formats or {}
    for ri in range(2, len(rows) + 2):
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.font = Font(name=F, size=10)
            cell.border = BOX
            cell.alignment = Alignment(vertical='top', wrap_text=ci in wrap_cols)
            if ci in input_cols:
                cell.fill = INPUT
            if ci in num_formats:
                cell.number_format = num_formats[ci]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)), len(rows) + 1)


def add_dropdown(ws, col, options, nrows):
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(options), allow_blank=True)
    dv.add('%s2:%s%d' % (col, col, nrows + 1))
    ws.add_data_validation(dv)


def summary_block(ws, start_row, title, labels, data_sheet, data_col, total_ref):
    ws.cell(row=start_row, column=1, value=title).font = Font(name=F, bold=True, size=11)
    hdr = start_row + 1
    for ci, h in enumerate(('Category', 'Count', '% of total'), 1):
        c = ws.cell(row=hdr, column=ci, value=h)
        c.font = Font(name=F, bold=True, size=10); c.fill = SUB; c.border = BOX
    r = hdr + 1
    first = r
    for lab in labels:
        ws.cell(row=r, column=1, value=lab).font = Font(name=F, size=10)
        ws.cell(row=r, column=2, value="=COUNTIF('%s'!$%s:$%s,A%d)" % (data_sheet, data_col, data_col, r))
        ws.cell(row=r, column=3, value='=IF(%s=0,0,B%d/%s)' % (total_ref, r, total_ref))
        for ci in (1, 2, 3):
            ws.cell(row=r, column=ci).border = BOX
            ws.cell(row=r, column=ci).font = Font(name=F, size=10)
        ws.cell(row=r, column=3).number_format = '0.0%'
        r += 1
    ws.cell(row=r, column=1, value='Total').font = Font(name=F, bold=True, size=10)
    ws.cell(row=r, column=2, value='=SUM(B%d:B%d)' % (first, r - 1)).font = Font(name=F, bold=True, size=10)
    for ci in (1, 2):
        ws.cell(row=r, column=ci).border = BOX
    return r + 2


def readme(ws, lines):
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 100
    for i, (a, b) in enumerate(lines, 1):
        ca, cb = ws.cell(row=i, column=1, value=a), ws.cell(row=i, column=2, value=b)
        heading = bool(a) and not b
        ca.font = Font(name=F, bold=True, size=12 if heading else 10)
        cb.font = Font(name=F, size=10)
        cb.alignment = Alignment(wrap_text=True, vertical='top')
        ca.alignment = Alignment(vertical='top', wrap_text=True)


# =============================== OVERLAPS ===================================
def build_overlaps():
    wb = Workbook()
    ws_s = wb.active; ws_s.title = 'Summary'
    ws = wb.create_sheet('Overlaps')
    ws_r = wb.create_sheet('How to read')

    items = sorted(R['overlaps'], key=lambda o: (task_label(o['tid'])[0], task_label(o['tid'])[1],
                                                 float(o['x']['t_start'])))
    rows = []
    for n, o in enumerate(items, 1):
        v, who, st = task_label(o['tid'])
        x, y = o['x'], o['y']
        rows.append([n, v, who, st, hand_name(o['hand']),
                     tc(x['t_start']), tc(x['t_end']), x['caption'],
                     tc(y['t_start']), tc(y['t_end']), y['caption'],
                     round(float(o['inter']), 2),
                     'Fully inside' if o['contained'] else 'Partial',
                     o['hint'], None, None])
    headers = ['No.', 'Video', 'Annotator', 'Task status', 'Hand',
               'Row 1 start', 'Row 1 end', 'Row 1 caption',
               'Row 2 start', 'Row 2 end', 'Row 2 caption',
               'Overlap (s)', 'Overlap type', 'Review hint', 'QA decision', 'QA notes']
    write_table(ws, headers, rows,
                [6, 17, 18, 10, 8, 10, 10, 44, 10, 10, 44, 10, 12, 40, 18, 30],
                wrap_cols=(8, 11, 14, 16), input_cols=(15, 16), num_formats={12: '0.00'})
    add_dropdown(ws, 'O', ['Correct as is', 'Fix row 1', 'Fix row 2', 'Delete row 1', 'Delete row 2', 'Other'], len(rows))

    s = ws_s
    s.column_dimensions['A'].width = 62; s.column_dimensions['B'].width = 12; s.column_dimensions['C'].width = 12
    s['A1'] = 'Same-hand overlaps - QA review'; s['A1'].font = Font(name=F, bold=True, size=14)
    s['A2'] = 'Generated %s from the live database. One row per overlapping pair.' % TODAY
    s['A2'].font = Font(name=F, size=10, italic=True)
    s['A4'] = 'Overlapping pairs'; s['A4'].font = Font(name=F, bold=True, size=11)
    s['B4'] = "=COUNTA(Overlaps!$A:$A)-1"; s['B4'].font = Font(name=F, bold=True, size=11)
    s['A5'] = 'Reviewed (QA decision filled in)'; s['A5'].font = Font(name=F, size=10)
    s['B5'] = "=COUNTA(Overlaps!$O:$O)-1"
    s['C5'] = '=IF(B4=0,0,B5/B4)'; s['C5'].number_format = '0.0%'
    for c in ('B5', 'C5'): s[c].font = Font(name=F, size=10)

    hints = [h for h, _ in collections.Counter(o['hint'] for o in items).most_common()]
    r = summary_block(s, 7, 'By review hint', hints, 'Overlaps', 'N', '$B$4')
    r = summary_block(s, r, 'By overlap type', ['Fully inside', 'Partial'], 'Overlaps', 'M', '$B$4')
    vids = [v for v, _ in collections.Counter(task_label(o['tid'])[0] for o in items).most_common()]
    r = summary_block(s, r, 'By video (most first)', vids, 'Overlaps', 'B', '$B$4')
    whos = [w for w, _ in collections.Counter(task_label(o['tid'])[1] for o in items).most_common()]
    summary_block(s, r, 'By annotator (most first)', whos, 'Overlaps', 'C', '$B$4')

    readme(ws_r, [
        ('What this file lists', ''),
        ('Overlap', 'Two annotations in the same task where the SAME hand has two different actions recorded '
                    'over the same moment. Example: "Right hand holds the pen" 30-35 s and "Right hand moves the '
                    'tweezers" 32-34 s.'),
        ('Not listed', 'A left-hand and a right-hand annotation at the same time. Both hands are annotated, so '
                       'that is expected.'),
        ('Not listed', 'Spans that only touch for less than one frame (1/30 s). Those are back-to-back stamps.'),
        ('Not listed', 'Pairs that are the same action recorded twice. Those are in the separate Duplicates file.'),
        ('Not checked', 'Captions that do not start with "Left hand", "Right hand" or "Both hands" - the hand '
                        'cannot be read. "Both hands" counts as left and right.'),
        ('', ''),
        ('Review hints', ''),
        ('Two different actions at once - likely error',
         'The same hand is doing two unrelated things at the same time. Usually one span is too long or one '
         'caption has the wrong hand.'),
        ('Holds one thing while doing something else - check',
         'One caption is "holds ..." and the other action is on a different object. Sometimes possible, often an error.'),
        ('Holds an object while acting on the same object - probably valid',
         'e.g. "holds the bottle" while "pours from the bottle". Usually fine; quick check.'),
        ('Same kind of object, different identifier',
         'e.g. tube A vs tube C, first vs second filter. One hand cannot normally act on both at once.'),
        ('Conflicting direction or side', 'e.g. picks up / places down, turns on / off, left / right shield at once.'),
        ('', ''),
        ('Columns', ''),
        ('Row 1 / Row 2', 'The two annotations. Row 1 starts first. Times are minutes:seconds within that video.'),
        ('Overlap (s)', 'How many seconds the two spans share.'),
        ('Overlap type', '"Fully inside" = one span lies completely within the other.'),
        ('QA decision / QA notes', 'Yellow cells are for reviewers. The Summary tab counts how many are done.'),
        ('', ''),
        ('Important', 'Review hints are automatic guidance from the caption text, not a verdict. The video decides.'),
    ])
    path = 'exports/INOUEE_QA_Overlaps_%s.xlsx' % TODAY
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path, len(rows)


# ============================== DUPLICATES ==================================
def build_duplicates():
    wb = Workbook()
    ws_s = wb.active; ws_s.title = 'Summary'
    ws_w = wb.create_sheet('Within-task duplicates')
    ws_c = wb.create_sheet('Cross-annotator duplicates')
    ws_r = wb.create_sheet('How to read')

    # --- within task
    items = sorted(R['dup_within'], key=lambda d: (task_label(d['tid'])[0], float(d['x']['t_start'])))
    rows = []
    for n, d in enumerate(items, 1):
        v, who, st = task_label(d['tid'])
        x, y = d['x'], d['y']
        rows.append([n, v, who, st, hand_name(d.get('hand', '?')), d['kind'],
                     tc(x['t_start']), tc(x['t_end']), x['caption'],
                     tc(y['t_start']), tc(y['t_end']), y['caption'],
                     round(float(d['iou']), 3), None, None])
    hw = ['No.', 'Video', 'Annotator', 'Task status', 'Hand', 'Duplicate type',
          'Row 1 start', 'Row 1 end', 'Row 1 caption', 'Row 2 start', 'Row 2 end', 'Row 2 caption',
          'Time overlap', 'QA decision', 'QA notes']
    write_table(ws_w, hw, rows, [6, 17, 18, 10, 8, 26, 10, 10, 42, 10, 10, 42, 10, 18, 30],
                wrap_cols=(9, 12, 15), input_cols=(14, 15), num_formats={13: '0%'})
    add_dropdown(ws_w, 'N', ['Keep row 1', 'Keep row 2', 'Keep both', 'Other'], len(rows))
    n_within = len(rows)

    # --- cross annotator
    citems = sorted(R['cross'], key=lambda c: (c['group'], float(c['x']['t_start'])))
    rows = []
    for n, c in enumerate(citems, 1):
        _, wa, sa = task_label(c['ta']); _, wb_, sb = task_label(c['tb'])
        x, y = c['x'], c['y']
        rows.append([n, c['group'], hand_name(c['hand']), c['verdict'],
                     wa, sa, tc(x['t_start']), tc(x['t_end']), x['caption'],
                     wb_, sb, tc(y['t_start']), tc(y['t_end']), y['caption'],
                     round(float(c['iou']), 3), None, None])
    hc = ['No.', 'Video', 'Hand', 'Agreement',
          'Annotator 1', 'Status 1', 'Start 1', 'End 1', 'Caption 1',
          'Annotator 2', 'Status 2', 'Start 2', 'End 2', 'Caption 2',
          'Time overlap', 'QA decision', 'QA notes']
    write_table(ws_c, hc, rows, [6, 22, 8, 20, 16, 9, 10, 10, 40, 16, 9, 10, 10, 40, 10, 18, 30],
                wrap_cols=(2, 9, 14, 17), input_cols=(16, 17), num_formats={15: '0%'})
    add_dropdown(ws_c, 'P', ['Keep annotator 1', 'Keep annotator 2', 'Merge', 'Other'], len(rows))
    n_cross = len(rows)

    # --- summary
    s = ws_s
    for col, w in (('A', 46), ('B', 18), ('C', 14), ('D', 18), ('E', 14), ('F', 16)):
        s.column_dimensions[col].width = w
    s['A1'] = 'Duplicates - QA review'; s['A1'].font = Font(name=F, bold=True, size=14)
    s['A2'] = 'Generated %s from the live database. One row per duplicated pair.' % TODAY
    s['A2'].font = Font(name=F, size=10, italic=True)
    s['A4'] = 'Within-task duplicate pairs'; s['B4'] = "=COUNTA('Within-task duplicates'!$A:$A)-1"
    s['A5'] = 'Cross-annotator duplicate pairs'; s['B5'] = "=COUNTA('Cross-annotator duplicates'!$A:$A)-1"
    s['A6'] = 'All duplicate pairs'; s['B6'] = '=B4+B5'
    for c in ('A4', 'A5', 'A6', 'B4', 'B5', 'B6'):
        s[c].font = Font(name=F, size=10, bold=c in ('A6', 'B6'))

    kinds = [k for k, _ in collections.Counter(d['kind'] for d in items).most_common()]
    r = summary_block(s, 8, 'Within-task duplicates, by type', kinds, 'Within-task duplicates', 'F', '$B$4')
    verds = [k for k, _ in collections.Counter(c['verdict'] for c in citems).most_common()]
    r = summary_block(s, r, 'Cross-annotator duplicates, by agreement', verds,
                      'Cross-annotator duplicates', 'D', '$B$5')

    s.cell(row=r, column=1, value='Videos annotated by more than one task').font = Font(name=F, bold=True, size=11)
    r += 1
    hdr = ['Video', 'Annotator 1', 'Rows 1', 'Annotator 2', 'Rows 2', 'Duplicated pairs']
    for ci, h in enumerate(hdr, 1):
        c = s.cell(row=r, column=ci, value=h)
        c.font = Font(name=F, bold=True, size=10); c.fill = SUB; c.border = BOX
    r += 1
    for st in R['cross_stats']:
        vals = [st['group'], task_label(st['ta'])[1], st['na'], task_label(st['tb'])[1], st['nb'],
                "=COUNTIF('Cross-annotator duplicates'!$B:$B,A%d)" % r]
        for ci, v in enumerate(vals, 1):
            c = s.cell(row=r, column=ci, value=v)
            c.font = Font(name=F, size=10); c.border = BOX
        r += 1
    s.cell(row=r, column=1, value='Rows 1 / Rows 2 are each task\'s full annotation count; '
                                  'duplicated pairs are the ones matched to the other task.').font = \
        Font(name=F, size=9, italic=True)

    readme(ws_r, [
        ('What this file lists', ''),
        ('Within-task duplicates', 'The same action recorded twice inside ONE annotator\'s task, same hand, '
                                   'overlapping in time.'),
        ('  Same caption recorded twice', 'Identical caption (ignoring capitals and spacing) and the two spans '
                                          'overlap by at least half. Almost always a double entry - delete one.'),
        ('  Same caption, spans partly overlap', 'Identical caption, but the spans overlap by less than half. '
                                                 'Could be a double entry, or one action whose end was stamped late.'),
        ('  Same action, reworded', 'Same verb (allowing synonyms such as grab / take / pick up), same object, '
                                    'overlapping by at least half, worded differently.'),
        ('', ''),
        ('Cross-annotator duplicates', 'Where the same footage was assigned to two tasks, both annotators '
                                       'annotated it. Each row pairs one annotation from each, same hand, '
                                       'overlapping by at least half. Each annotation is used in at most one pair.'),
        ('  Agree - same wording', 'Both wrote the same caption.'),
        ('  Agree - reworded', 'Same action and object, different words.'),
        ('  Unclear - check video', 'The captions describe a different action or object, OR say the same thing '
                                    'in words the automatic check cannot match (e.g. "opens" vs "lifts" the '
                                    'filter window). Watch the clip to decide.'),
        ('', ''),
        ('E10-S01 (original + re-upload)', 'The first upload of E10-S01 was corrupt and was re-uploaded as '
                                          '"E10-S01 new". The original task still holds 251 annotations for the '
                                          'first ~10 minutes, which the same annotator then redid on the new '
                                          'file. These pairs show that overlap. Recommend archiving the '
                                          'original task rather than reviewing each row.'),
        ('', ''),
        ('Not duplicates', 'The same caption at different times (an action repeated later) is normal and is '
                           'not listed. Different objects (tube A vs tube C) are never treated as duplicates - '
                           'those overlaps are in the Overlaps file.'),
        ('Time overlap', 'Shared time divided by the combined span of the pair: 100% = identical timing.'),
        ('QA decision / QA notes', 'Yellow cells are for reviewers.'),
    ])
    path = 'exports/INOUEE_QA_Duplicates_%s.xlsx' % TODAY
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path, n_within, n_cross


p1, n1 = build_overlaps()
p2, n2, n3 = build_duplicates()
print(p1, n1)
print(p2, n2, n3)
